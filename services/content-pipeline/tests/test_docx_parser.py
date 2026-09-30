from __future__ import annotations

import io

import docx

from content_pipeline.parsers.docx import parse_docx


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Capaian Pembelajaran", level=1)
    document.add_paragraph("Peserta didik dapat membandingkan bilangan.")
    document.add_heading("Fase A", level=2)
    document.add_paragraph("Kelas I dan II.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Fase"
    table.rows[0].cells[1].text = "Kelas"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_sections_follow_heading_styles() -> None:
    parsed = parse_docx(_docx_bytes())
    assert parsed.document_type == "docx"
    assert [s.heading for s in parsed.sections] == ["Capaian Pembelajaran", "Fase A"]
    assert [s.level for s in parsed.sections] == [1, 2]


def test_docx_includes_tables() -> None:
    parsed = parse_docx(_docx_bytes())
    assert "Fase | Kelas" in parsed.sections[1].text
